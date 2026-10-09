#!/usr/bin/env python3
"""Export the laptop workbook (xlsx) to the JSON files the API loads.

The spreadsheet stays the place where data is edited (it is friendly for
filling in specs and prices). This script turns it into clean, typed JSON so
the backend never has to parse Excel at runtime.

    python scripts/export_data.py                # write JSON into the API package
    python scripts/export_data.py --check        # fail if the committed JSON is stale

Rules it follows:
  * Example rows ("contoh-...", "(CONTOH)") are skipped.
  * "unknown" or blank in a numeric column becomes null. Unknown stays unknown;
    nothing is ever guessed.
  * The workbook layout is verified; if someone moves a column or a section the
    script stops with a readable message instead of exporting wrong data.
  * Cached cell values are used, so save the workbook in Excel/LibreOffice
    (which recalculates formulas) before exporting.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = ROOT / "data" / "source" / "data-laptop.xlsx"
DEFAULT_OUT = ROOT / "apps" / "api" / "laptop_advisor" / "data"

# --------------------------------------------------------------------------- #
# Column maps: JSON key -> (workbook key in header row 3, kind)
# kind: str | num | date
# --------------------------------------------------------------------------- #
MODEL_FIELDS: dict[str, tuple[str, str]] = {
    "id": ("kunci_model", "str"),
    "brand": ("merek", "str"),
    "series": ("seri", "str"),
    "year": ("tahun_rilis", "num"),
    "sales_status": ("status", "str"),
    "body_material": ("bahan", "str"),
    "screen_inches": ("layar_inci", "num"),
    "usb_a_count": ("usb_a", "num"),
    "usb_c_count": ("usb_c", "num"),
    "usb_c_kind": ("usbc_jenis", "str"),
    "usb_c_charging": ("usbc_charge", "str"),
    "hdmi": ("hdmi", "str"),
    "hdmi_version": ("hdmi_versi", "str"),
    "sd_card": ("sd", "str"),
    "lan": ("lan", "str"),
    "summary": ("ringkasan", "str"),
    "watch_out": ("waspada", "str"),
    "common_issues": ("masalah_umum", "str"),
    "spec_url": ("url_spek", "str"),
    "checked_on": ("tgl_verif", "date"),
}

VARIANT_FIELDS: dict[str, tuple[str, str]] = {
    "id": ("id_varian", "str"),
    "model_id": ("kunci_model", "str"),
    "code": ("kode_varian", "str"),
    "year": ("tahun_rilis", "num"),
    "sales_status": ("status_penjualan", "str"),
    "discontinued_year": ("tahun_dihentikan", "num"),
    "market": ("pasar", "str"),
    "official_segment": ("segmen_resmi", "str"),
    "official_target": ("target_resmi", "str"),
    "cpu": ("cpu", "str"),
    "gpu": ("gpu", "str"),
    "gpu_tgp_w": ("gpu_tgp", "num"),
    "ram_gb": ("ram_gb", "num"),
    "ram_type": ("ram_tipe", "str"),
    "ram_speed_mts": ("ram_mts", "num"),
    "ram_channel": ("ram_channel", "str"),
    "ram_layout": ("ram_layout", "str"),
    "ram_slots_total": ("ram_slot_total", "num"),
    "ram_slots_free": ("ram_slot_kosong", "num"),
    "ram_soldered_gb": ("ram_disolder_gb", "num"),
    "ram_max_gb": ("ram_maks_gb", "num"),
    "storage_gb": ("st_gb", "num"),
    "storage_type": ("st_tipe", "str"),
    "storage_interface": ("st_antarmuka", "str"),
    "storage_form": ("st_form", "str"),
    "storage_replaceable": ("st_bisa_diganti", "str"),
    "m2_slots_total": ("st_m2_total", "num"),
    "m2_slots_free": ("st_m2_kosong", "num"),
    "bay_25_free": ("st_25_kosong", "str"),
    "upgrade_verification": ("verif_upgrade", "str"),
    "upgrade_source": ("sumber_upgrade", "str"),
    "weight_kg": ("bobot_kg", "num"),
    "panel": ("panel", "str"),
    "resolution": ("resolusi", "str"),
    "refresh_hz": ("refresh_hz", "num"),
    "touch": ("sentuh", "str"),
    "battery_wh": ("baterai_wh", "num"),
    "charger_w": ("charger_w", "num"),
    "wifi": ("wifi", "str"),
    "bluetooth": ("bluetooth", "str"),
    "fingerprint": ("sidik_jari", "str"),
    "backlit_keyboard": ("kb_lampu", "str"),
    "webcam": ("webcam", "str"),
    "os": ("os", "str"),
    "os_edition": ("os_edisi", "str"),
    "office": ("office", "str"),
    "warranty_type": ("garansi_jenis", "str"),
    "warranty_months": ("garansi_bulan", "num"),
    "battery_warranty_months": ("garansi_baterai", "num"),
    "available_new": ("tersedia_baru", "str"),
    "available_used": ("tersedia_bekas", "str"),
    "editor_note": ("catatan_editor", "str"),
}

CPU_FIELDS: dict[str, tuple[str, str]] = {
    "name": ("nama", "str"),
    "maker": ("produsen", "str"),
    "generation": ("generasi", "str"),
    "codename": ("kode", "str"),
    "year": ("tahun", "num"),
    "cores": ("core", "num"),
    "threads": ("thread", "num"),
    "power_class": ("kelas", "str"),
    "windows11": ("win11", "str"),
    "tier": ("tier", "num"),
    "tier_reason": ("alasan", "str"),
    "plain": ("awam", "str"),
    "source": ("sumber", "str"),
    "checked_on": ("tgl", "date"),
}

GPU_FIELDS: dict[str, tuple[str, str]] = {
    "name": ("nama", "str"),
    "maker": ("produsen", "str"),
    "kind": ("jenis", "str"),
    "architecture": ("arsitektur", "str"),
    "year": ("tahun", "num"),
    "vram_gb": ("vram", "num"),
    "tier": ("tier", "num"),
    "tier_reason": ("alasan", "str"),
    "plain": ("awam", "str"),
    "source": ("sumber", "str"),
    "checked_on": ("tgl", "date"),
}

PRICE_FIELDS: dict[str, tuple[str, str]] = {
    "variant_id": ("id_varian", "str"),
    "condition": ("kondisi", "str"),
    "low": ("hmin", "num"),
    "high": ("hmaks", "num"),
    "mid": ("tengah", "num"),
    "source_type": ("jenis_sumber", "str"),
    "source": ("sumber", "str"),
    "checked_on": ("tgl", "date"),
    "note": ("catatan", "str"),
}

# The six criteria, in the fixed order used by the Aturan sheet.
CRITERIA = [
    ("cpu", "Prosesor"),
    ("gpu", "Kartu grafis"),
    ("ram", "RAM"),
    ("storage", "Penyimpanan"),
    ("display", "Jenis layar"),
    ("refresh", "Refresh rate"),
]


class ExportError(Exception):
    """The workbook does not look the way this script expects."""


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #
def is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def is_example(v: Any) -> bool:
    if not isinstance(v, str):
        return False
    low = v.strip().lower()
    return low.startswith("contoh") or "(contoh)" in low


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    if not slug:
        raise ExportError(f"Nama '{text}' tidak bisa dijadikan alamat (slug).")
    return slug


def number(x: float | int) -> int | float:
    return int(x) if float(x).is_integer() else float(x)


def convert(value: Any, kind: str, where: str, issues: list[str]) -> Any:
    if value is None:
        return None
    if kind == "str":
        text = str(value).strip()
        return text or None
    if kind == "num":
        if isinstance(value, bool):
            issues.append(f"{where}: nilai ya/tidak dibaca sebagai angka, dikosongkan.")
            return None
        if isinstance(value, (int, float)):
            return number(value)
        text = str(value).strip()
        if text.lower() in {"", "unknown", "-"}:
            return None
        try:
            return number(float(text.replace(",", ".")))
        except ValueError:
            issues.append(f"{where}: '{text}' bukan angka, dikosongkan.")
            return None
    if kind == "date":
        if isinstance(value, dt.datetime):
            return value.date().isoformat()
        if isinstance(value, dt.date):
            return value.isoformat()
        text = str(value).strip()
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return text
        if text.lower() in {"", "unknown", "-"}:
            return None
        issues.append(f"{where}: '{text}' bukan tanggal (YYYY-MM-DD), dikosongkan.")
        return None
    raise AssertionError(kind)


def header_keys(ws, row: int = 3) -> dict[str, int]:
    """Header cells look like 'Label\\nkey'; the key is the stable part."""
    keys: dict[str, int] = {}
    for cell in ws[row]:
        if isinstance(cell.value, str) and "\n" in cell.value:
            keys[cell.value.split("\n")[-1].strip()] = cell.column
    return keys


def read_table(ws, fields: dict[str, tuple[str, str]], id_field: str, issues: list[str],
               first_row: int = 4) -> list[dict[str, Any]]:
    keys = header_keys(ws)
    missing = sorted({sk for sk, _ in fields.values() if sk not in keys})
    if missing:
        raise ExportError(f"Tab {ws.title}: kolom tidak ditemukan di baris 3: {', '.join(missing)}")
    id_col = keys[fields[id_field][0]]
    rows: list[dict[str, Any]] = []
    for r in range(first_row, ws.max_row + 1):
        raw_id = ws.cell(r, id_col).value
        if is_blank(raw_id) or is_example(raw_id):
            continue
        rec: dict[str, Any] = {}
        for jkey, (skey, kind) in fields.items():
            rec[jkey] = convert(ws.cell(r, keys[skey]).value, kind, f"{ws.title} baris {r} kolom {skey}", issues)
        if any(is_example(v) for v in rec.values() if isinstance(v, str)):
            continue
        rows.append(rec)
    return rows


# --------------------------------------------------------------------------- #
# Aturan sheet
# --------------------------------------------------------------------------- #
def find_sections(ws) -> dict[int, int]:
    found: dict[int, int] = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str):
            m = re.match(r"^(\d)\.\s", v)
            if m:
                found[int(m.group(1))] = r
    absent = [n for n in range(1, 8) if n not in found]
    if absent:
        raise ExportError(f"Tab Aturan: bagian bernomor {absent} tidak ditemukan.")
    return found


def read_rules(wb, issues: list[str]) -> dict[str, Any]:
    ws = wb["Aturan"]
    sec = find_sections(ws)
    cell = lambda r, c: ws.cell(r, c).value  # noqa: E731

    def text(r: int, c: int) -> str | None:
        return convert(cell(r, c), "str", f"Aturan {r}:{c}", issues)

    # 1. status and verdict meanings (two blocks separated by a blank row)
    r = sec[1] + 2
    statuses = []
    while not is_blank(cell(r, 1)):
        statuses.append({"label": text(r, 1), "meaning": text(r, 2)})
        r += 1
    while is_blank(cell(r, 1)):
        r += 1
    r += 1  # header row of the verdict block
    verdicts = []
    while not is_blank(cell(r, 1)):
        verdicts.append({"label": text(r, 1), "meaning": text(r, 2)})
        r += 1

    # 2. needs
    r = sec[2] + 2
    needs = []
    while not is_blank(cell(r, 1)):
        name = text(r, 1)
        needs.append({"id": slugify(name), "name": name, "description": text(r, 2)})
        r += 1
    if len({n["id"] for n in needs}) != len(needs):
        raise ExportError("Tab Aturan: dua kebutuhan menghasilkan alamat (slug) yang sama.")

    # 3. thresholds: 6 rows per need, fixed criterion order
    rule0 = sec[3] + 3
    thresholds: dict[str, list[dict[str, Any]]] = {}
    criteria_meta: list[dict[str, Any]] = []
    for j, need in enumerate(needs):
        entries = []
        for k, (ckey, prefix) in enumerate(CRITERIA):
            row = rule0 + len(CRITERIA) * j + k
            owner, crit_name = cell(row, 1), cell(row, 2)
            if owner != need["name"]:
                raise ExportError(f"Tab Aturan baris {row}: kebutuhan '{owner}' tidak sesuai urutan (seharusnya '{need['name']}').")
            if not isinstance(crit_name, str) or not crit_name.startswith(prefix):
                raise ExportError(f"Tab Aturan baris {row}: kriteria '{crit_name}' tidak sesuai urutan (seharusnya diawali '{prefix}').")
            entries.append({
                "criterion": ckey,
                "minimum": convert(cell(row, 3), "num", f"Aturan C{row}", issues),
                "comfortable": convert(cell(row, 4), "num", f"Aturan D{row}", issues),
                "reason": text(row, 6),
                "basis": text(row, 7),
            })
            if j == 0:
                criteria_meta.append({"key": ckey, "name": crit_name, "unit": text(row, 5)})
        thresholds[need["id"]] = entries

    # 4. display panel ranks
    r = sec[4] + 2
    panels = []
    while not is_blank(cell(r, 1)) and isinstance(cell(r, 2), (int, float)):
        panels.append({"name": text(r, 1), "rank": number(cell(r, 2)), "note": text(r, 3)})
        r += 1
    if not panels:
        raise ExportError("Tab Aturan: tabel peringkat jenis layar kosong.")

    # 5. general warnings
    r = sec[5] + 2
    warnings = []
    while isinstance(cell(r, 1), str) and re.fullmatch(r"P\d+", cell(r, 1).strip()):
        warnings.append({"code": text(r, 1), "when": text(r, 2), "text": text(r, 5)})
        r += 1
    if [w["code"] for w in warnings] != [f"P{i}" for i in range(1, 7)]:
        raise ExportError("Tab Aturan: peringatan umum harus P1 sampai P6 berurutan.")

    # 6. tier meanings
    r = sec[6] + 2
    tiers = []
    while isinstance(cell(r, 1), (int, float)):
        tiers.append({"tier": number(cell(r, 1)), "cpu": text(r, 2), "gpu": text(r, 6)})
        r += 1

    # 7. not scored yet
    r = sec[7] + 2
    not_scored = []
    while not is_blank(cell(r, 1)):
        not_scored.append({"item": text(r, 1), "reason": text(r, 2)})
        r += 1

    return {
        "needs": needs,
        "criteria": criteria_meta,
        "thresholds": thresholds,
        "panels": panels,
        "warnings": warnings,
        "tiers": tiers,
        "not_scored": not_scored,
        "statuses": statuses,
        "verdicts": verdicts,
    }


# --------------------------------------------------------------------------- #
# Cross checks
# --------------------------------------------------------------------------- #
def cross_check(models, variants, cpus, gpus, prices, issues: list[str]) -> None:
    def dupes(items: list[str], what: str) -> None:
        seen: set[str] = set()
        for it in items:
            low = it.lower()
            if low in seen:
                issues.append(f"{what} ganda: {it}")
            seen.add(low)

    dupes([m["id"] for m in models], "Kunci model")
    dupes([v["id"] for v in variants], "ID varian")
    dupes([c["name"] for c in cpus], "Nama CPU")
    dupes([g["name"] for g in gpus], "Nama GPU")

    model_ids = {m["id"].lower() for m in models}
    cpu_names = {c["name"].lower() for c in cpus}
    gpu_names = {g["name"].lower() for g in gpus}
    variant_ids = {v["id"].lower() for v in variants}

    for v in variants:
        if not v["model_id"] or v["model_id"].lower() not in model_ids:
            issues.append(f"Varian {v['id']}: model '{v['model_id']}' tidak ada di tab Model.")
        if v["cpu"] and v["cpu"].lower() not in cpu_names:
            issues.append(f"Varian {v['id']}: CPU '{v['cpu']}' tidak ada di tab CPU.")
        if v["gpu"] and v["gpu"].lower() not in gpu_names:
            issues.append(f"Varian {v['id']}: GPU '{v['gpu']}' tidak ada di tab GPU.")
    for c in cpus:
        if c["tier"] is not None and c["tier"] not in (1, 2, 3, 4, 5):
            issues.append(f"CPU {c['name']}: tingkat {c['tier']} di luar 1-5.")
    for g in gpus:
        if g["tier"] is not None and g["tier"] not in (1, 2, 3, 4, 5):
            issues.append(f"GPU {g['name']}: tingkat {g['tier']} di luar 1-5.")
    for p in prices:
        if not p["variant_id"] or p["variant_id"].lower() not in variant_ids:
            issues.append(f"Harga untuk '{p['variant_id']}': varian tidak ada di tab Varian.")
        if p["condition"] not in ("baru", "bekas"):
            issues.append(f"Harga {p['variant_id']}: kondisi '{p['condition']}' harus baru atau bekas.")
        if p["low"] is None and p["high"] is None:
            issues.append(f"Harga {p['variant_id']}: harga terendah dan tertinggi sama-sama kosong.")
        if p["low"] is not None and p["high"] is not None and p["low"] > p["high"]:
            issues.append(f"Harga {p['variant_id']}: harga terendah lebih besar dari harga tertinggi.")
        if not p["checked_on"]:
            issues.append(f"Harga {p['variant_id']}: tanggal dicek kosong.")


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def build(source: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(source, data_only=True)
    for sheet in ("Model", "Varian", "CPU", "GPU", "Harga", "Aturan"):
        if sheet not in wb.sheetnames:
            raise ExportError(f"Tab '{sheet}' tidak ditemukan di {source.name}.")
    issues: list[str] = []
    models = read_table(wb["Model"], MODEL_FIELDS, "id", issues)
    variants = read_table(wb["Varian"], VARIANT_FIELDS, "id", issues)
    cpus = read_table(wb["CPU"], CPU_FIELDS, "name", issues)
    gpus = read_table(wb["GPU"], GPU_FIELDS, "name", issues)
    prices = read_table(wb["Harga"], PRICE_FIELDS, "variant_id", issues)
    for p in prices:
        if p["mid"] is None and (p["low"] is not None or p["high"] is not None):
            vals = [x for x in (p["low"], p["high"]) if x is not None]
            p["mid"] = number(sum(vals) / len(vals))
    rules = read_rules(wb, issues)
    cross_check(models, variants, cpus, gpus, prices, issues)

    dates = [d for d in [p["checked_on"] for p in prices] + [m["checked_on"] for m in models] if d]
    meta = {
        "source_file": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "generated_at": dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
        "last_checked": max(dates) if dates else None,
        "counts": {
            "models": len(models), "variants": len(variants), "cpus": len(cpus),
            "gpus": len(gpus), "prices": len(prices), "needs": len(rules["needs"]),
        },
        "issues": issues,
    }
    return {"meta": meta, "rules": rules, "models": models, "variants": variants,
            "cpus": cpus, "gpus": gpus, "prices": prices}


def dump(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=2) + "\n"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", type=Path, default=DEFAULT_SOURCE, help="path to the workbook")
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT, help="folder for the JSON files")
    ap.add_argument("--check", action="store_true", help="only verify that the JSON in --out is up to date")
    args = ap.parse_args(argv)

    try:
        data = build(args.source)
    except ExportError as err:
        print(f"GAGAL: {err}", file=sys.stderr)
        return 2

    files = {f"{name}.json": dump(data[name]) for name in ("meta", "rules", "models", "variants", "cpus", "gpus", "prices")}
    meta = data["meta"]

    if args.check:
        stale = []
        for fname, content in files.items():
            path = args.out / fname
            if not path.exists():
                stale.append(f"{fname} (belum ada)")
                continue
            if fname == "meta.json":
                old = json.loads(path.read_text(encoding="utf-8"))
                new = json.loads(content)
                for volatile in ("generated_at",):
                    old.pop(volatile, None)
                    new.pop(volatile, None)
                if old != new:
                    stale.append(fname)
            elif path.read_text(encoding="utf-8") != content:
                stale.append(fname)
        if stale:
            print("JSON sudah kedaluwarsa terhadap spreadsheet: " + ", ".join(stale), file=sys.stderr)
            print("Jalankan: python scripts/export_data.py", file=sys.stderr)
            return 1
        print("JSON sudah sesuai dengan spreadsheet.")
        return 0

    args.out.mkdir(parents=True, exist_ok=True)
    for fname, content in files.items():
        (args.out / fname).write_text(content, encoding="utf-8")
    c = meta["counts"]
    print(f"Ekspor selesai ke {args.out}")
    print(f"  {c['models']} model, {c['variants']} varian, {c['cpus']} CPU, {c['gpus']} GPU, "
          f"{c['prices']} pengamatan harga, {c['needs']} kebutuhan; data dicek terakhir {meta['last_checked']}")
    if meta["issues"]:
        print(f"  {len(meta['issues'])} catatan data perlu dicek:")
        for issue in meta["issues"]:
            print(f"   - {issue}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
